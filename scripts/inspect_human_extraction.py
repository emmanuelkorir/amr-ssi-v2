"""Dump the human data-extraction workbook (values + comments) for review."""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "amr-ssis-ecsa-files" / "amr-ssis-ecsa" / "data-extraction" / "data-extraction.xlsx"


def cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ⏎ ")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for ws in wb.worksheets:
        print(f"\n{'=' * 100}\nSHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}\n{'=' * 100}")
        for row in ws.iter_rows():
            vals = [(c.coordinate, cell_text(c.value)) for c in row if c.value not in (None, "")]
            if vals:
                print(" | ".join(f"{coord}={val}" for coord, val in vals))
            for c in row:
                if c.comment:
                    print(f"  COMMENT {c.coordinate}: {c.comment.text}")


if __name__ == "__main__":
    main()
