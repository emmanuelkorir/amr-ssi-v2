"""Build the validated CSV/XLSX extraction dataset.

Every evidence id is resolved to a verbatim page slice from corpus/json and
asserted to be a substring of that page.  Numeric fields are re-checked
against their denominators; derived values and arithmetic mismatches are
flagged rather than silently corrected.
"""
from __future__ import annotations

import csv
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import records as R  # noqa: E402
from evidence_lib import page_text, resolve_quote  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
TOL = 0.5  # percentage-point tolerance for reported vs recomputed values


def pct(n: int, d: int) -> float | None:
    return round(100.0 * n / d, 2) if d else None


def pct_check(n: int, d: int, reported) -> tuple[float | None, str]:
    calc = pct(n, d)
    if calc is None:
        return None, "NO DENOMINATOR"
    if reported is None:
        return calc, "OK (not reported)"
    if abs(calc - float(reported)) <= TOL:
        return calc, "OK"
    return calc, f"MISMATCH (reported {reported}, calculated {calc})"


# ---------------------------------------------------------------------------
# 1. Resolve evidence
# ---------------------------------------------------------------------------
errors: list[str] = []
evidence: dict[str, dict] = {}
for spec in R.EV_SPECS:
    eid = spec["evidence_id"]
    if eid in evidence:
        errors.append(f"duplicate evidence id: {eid}")
    try:
        quote = resolve_quote(spec["doc_id"], spec["page"], spec["start"], spec["end"])
    except Exception as exc:  # noqa: BLE001
        errors.append(f"{eid}: {exc}")
        continue
    page = page_text(spec["doc_id"], spec["page"])
    if quote not in page:
        errors.append(f"{eid}: resolved quote is not a verbatim page substring")
    evidence[eid] = {**spec, "quote": quote}


def ev_cols(eid: str | None, suffix: str = "") -> dict:
    if not eid:
        return {f"evidence_id{suffix}": "", f"evidence_page{suffix}": "", f"evidence_quote{suffix}": ""}
    e = evidence.get(eid)
    if e is None:
        errors.append(f"unresolved evidence reference: {eid}")
        return {f"evidence_id{suffix}": eid, f"evidence_page{suffix}": "", f"evidence_quote{suffix}": ""}
    return {
        f"evidence_id{suffix}": eid,
        f"evidence_page{suffix}": e["page"],
        f"evidence_quote{suffix}": e["quote"],
    }


# ---------------------------------------------------------------------------
# 2. Incidence table
# ---------------------------------------------------------------------------
inc_rows = []
for row in R.INCIDENCE:
    n, d = int(row["n_ssi"]), int(row["n_total"])
    if n > d:
        errors.append(f"incidence numerator > denominator: {row['study_id']} {n}/{d}")
    if n < 0 or d <= 0:
        errors.append(f"invalid incidence counts: {row['study_id']} {n}/{d}")
    calc, check = pct_check(n, d, row.get("pct_reported"))
    eid2 = R.SECOND_EVIDENCE.get((row["study_id"], row["cohort_label"]))
    inc_rows.append({
        "study_id": row["study_id"],
        "cohort_label": row["cohort_label"],
        "n_ssi_events": n,
        "n_total": d,
        "pct_reported": row.get("pct_reported"),
        "pct_calculated": calc,
        "percent_check": check,
        "estimate_type": row["estimate_type"],
        "numerator_definition": row["numerator_definition"],
        "denominator_definition": row["denominator_definition"],
        "derived_flag": row["derived_flag"],
        "include_primary": row["include_primary"],
        "notes": row["notes"],
        **ev_cols(row["evidence_id"]),
        **ev_cols(eid2, "_2"),
    })
incidence = pd.DataFrame(inc_rows)

# ---------------------------------------------------------------------------
# 3. AMR table
# ---------------------------------------------------------------------------
amr_rows = []
for row in R.AMR:
    n, d = int(row["n_resistant"]), int(row["n_tested"])
    if n > d:
        errors.append(f"AMR numerator > denominator: {row['study_id']} {n}/{d}")
    calc, check = pct_check(n, d, row.get("pct_reported"))
    amr_rows.append({
        "study_id": row["study_id"],
        "pathogen": row["pathogen"],
        "pathogen_group": row["pathogen_group"],
        "resistance_outcome": row["resistance_outcome"],
        "resistance_marker": row["marker"],
        "n_resistant": n,
        "n_tested": d,
        "pct_reported": row.get("pct_reported"),
        "pct_calculated": calc,
        "percent_check": check,
        "denominator_source": row["denominator_source"],
        "derived_flag": row["derived_flag"],
        "include_primary": row["include_primary"],
        "notes": row["notes"],
        **ev_cols(row["evidence_id"]),
        **ev_cols(row.get("evidence_id_2"), "_2"),
    })
amr = pd.DataFrame(amr_rows)

# ---------------------------------------------------------------------------
# 4. Studies, outcomes, drivers, quality, discrepancies, screening
# ---------------------------------------------------------------------------
def with_evidence(rows: list[dict]) -> pd.DataFrame:
    out = []
    for row in rows:
        out.append({**row, **ev_cols(row.get("evidence_id"))})
    return pd.DataFrame(out)


EXCLUDED_LABELS = {
    "aaron-oladipo-et-al": "FALCON secondary analysis (Aaron Oladipo 2024)",
    "copy-of-gashaw-et-al-amr": "Copy of Gashaw 2026",
    "copy-of-mbabazi-et-al-amr": "Copy of Mbabazi 2025",
    "data-extraction": "Human data-extraction workbook",
    "kihla-et-al": "Kihla 2014 (mixed wounds)",
    "kumburu-et-al": "Kumburu 2015 (all hospital infections)",
    "laura-jung-et-al": "Jung 2024 (post-traumatic wounds)",
    "nardo-et-al": "De Nardo 2016 (duplicate cohort)",
    "ojulong-et-al": "Ojulong 2009 (isolate-level MRSA)",
    "ojulongandezera": "Ojulong 2010 (duplicate cohort)",
    "pmc10364890-clean-cut-ethiopia": "Clean Cut pilot duplicate",
    "protocol": "Review protocol",
    "response-to-reviewers-1": "Response to reviewers",
    "similar-sr-to-ssis-in-ecsa": "Similar systematic review correction",
    "soi230088-featured-1707866555-17813": "One-page image",
    "surgical-site-infection-among-patients-undergone-orthopaedic-surgery-at-muhimbil": "Orthopaedic DOCX duplicate",
    "worku-et-al-molecular": "Worku 2023 molecular report",
    "worku-et-al-b5820": "Worku 2023 molecular report",
    "worku-et-al-b6295820": "Worku 2023 molecular report",
}

studies = pd.DataFrame(R.STUDIES)
outcomes = with_evidence(R.OUTCOMES)
drivers = with_evidence(R.DRIVERS)
quality = with_evidence(R.QUALITY)
discrepancies = with_evidence(R.DISCREPANCIES)
screening = pd.DataFrame(R.SCREENING)
screening.loc[screening["decision"] == "Excluded", "label"] = screening.loc[
    screening["decision"] == "Excluded", "doc_id"].map(EXCLUDED_LABELS).fillna(
    screening.loc[screening["decision"] == "Excluded", "doc_id"])

# Enrich study table with source document and extraction status
studies["source_doc_id"] = studies["study_id"].map(R.PRIMARY_DOC)
studies["has_incidence"] = studies["study_id"].isin(incidence.loc[incidence["include_primary"] == "yes", "study_id"])
studies["has_amr"] = studies["study_id"].isin(amr["study_id"])
studies["n_incidence_rows"] = studies["study_id"].map(incidence.groupby("study_id").size()).fillna(0).astype(int)
studies["n_primary_incidence_rows"] = studies["study_id"].map(
    incidence.loc[incidence["include_primary"] == "yes"].groupby("study_id").size()).fillna(0).astype(int)
studies["n_amr_rows"] = studies["study_id"].map(amr.groupby("study_id").size()).fillna(0).astype(int)

# Attach study metadata and priority tags to the analysis tables
_meta = ["study_id", "first_author", "publication_year", "country", "subregion",
         "hospital_level", "design", "specialty"]
for _df in (incidence, amr):
    _add = studies[_meta].drop(columns="study_id")
    for _c in _add.columns:
        _df[_c] = _df["study_id"].map(studies.set_index("study_id")[_c])

_priority_map = {
    "MRSA": "Priority 1: MRSA",
    "3GC resistance": "Priority 2: 3GC/ESBL",
    "3GC non-susceptibility": "Priority 2: 3GC/ESBL",
    "ESBL": "Priority 2: 3GC/ESBL",
    "Carbapenem resistance": "Priority 3: carbapenem",
    "Carbapenemase production": "Supplementary: carbapenemase",
    "MDR": "Supplementary: MDR",
    "MR-CNS": "Supplementary: MR-CNS",
    "Any AMR": "Supplementary: any AMR",
}
amr["priority_outcome"] = amr["resistance_outcome"].map(_priority_map)

# ---------------------------------------------------------------------------
# 5. Structural validation
# ---------------------------------------------------------------------------
# every study has at least one row
for sid in studies["study_id"]:
    if sid not in set(incidence["study_id"]) | set(amr["study_id"]) | set(outcomes["study_id"]) | set(drivers["study_id"]):
        errors.append(f"study has no extracted outcome/AMR/driver row: {sid}")

# overlap groups: at most one primary incidence cohort unless explicitly distinct
overlap = studies[studies["overlap_group"].fillna("") != ""].groupby("overlap_group")
for group, g in overlap:
    primary = set(incidence.loc[(incidence["study_id"].isin(g["study_id"])) &
                                (incidence["include_primary"] == "yes"), "study_id"])
    if len(primary) > 1 and group != "clean-cut-ethiopia":
        errors.append(f"overlap group {group} contributes multiple primary cohorts: {sorted(primary)}")

# every screening doc has a decision and all corpus docs are logged
manifest_docs = {r["doc_id"] for r in csv.DictReader((ROOT / "corpus" / "manifest.csv").open())}
if set(screening["doc_id"]) != manifest_docs:
    errors.append(f"screening log does not match manifest: missing={manifest_docs - set(screening['doc_id'])}, extra={set(screening['doc_id']) - manifest_docs}")

# evidence ids referenced but not defined
refs = set()
for df, cols in [(incidence, ["evidence_id", "evidence_id_2"]), (amr, ["evidence_id", "evidence_id_2"]),
                 (outcomes, ["evidence_id"]), (drivers, ["evidence_id"]), (quality, ["evidence_id"]),
                 (discrepancies, ["evidence_id"])]:
    for c in cols:
        refs |= {v for v in df[c].dropna().astype(str) if v}
undefined = refs - set(evidence)
if undefined:
    errors.append(f"undefined evidence ids: {sorted(undefined)}")

# page columns should export as integers, not 2.0
for _df in (incidence, amr, outcomes, drivers, quality, discrepancies):
    for _c in ("evidence_page", "evidence_page_2"):
        if _c in _df.columns:
            _df[_c] = pd.to_numeric(_df[_c], errors="coerce").astype("Int64")

# ---------------------------------------------------------------------------
# 6. Write CSV
# ---------------------------------------------------------------------------
tables = {
    "studies": studies,
    "incidence": incidence,
    "amr": amr,
    "outcomes": outcomes,
    "drivers": drivers,
    "quality": quality,
    "discrepancies": discrepancies,
    "screening": screening,
    "evidence": pd.DataFrame([
        {k: v for k, v in e.items() if k not in {"start", "end"}} for e in R.EV_SPECS
    ]).assign(quote=lambda d: d["evidence_id"].map(lambda x: evidence.get(x, {}).get("quote", ""))),
}
for name, df in tables.items():
    df.to_csv(OUT / f"{name}.csv", index=False, encoding="utf-8-sig")

# ---------------------------------------------------------------------------
# 7. Validation report
# ---------------------------------------------------------------------------
mismatch_inc = incidence[incidence["percent_check"].str.startswith("MISMATCH")]
mismatch_amr = amr[amr["percent_check"].str.startswith("MISMATCH")]
derived_inc = incidence[incidence["derived_flag"] == "derived"]
derived_amr = amr[amr["derived_flag"] == "derived"]
primary_inc = incidence[incidence["include_primary"] == "yes"]

report = []
report.append("ECSA SSI/AMR extraction - validation report")
report.append("=" * 55)
report.append(f"Studies: {len(studies)} (incidence component {int(studies['incidence_component'].sum())}, AMR component {int(studies['amr_component'].sum())})")
report.append(f"Primary incidence rows: {len(primary_inc)} covering {primary_inc['study_id'].nunique()} cohorts; all incidence rows: {len(incidence)}")
report.append(f"AMR rows: {len(amr)} covering {amr['study_id'].nunique()} studies")
no_amr_rows = sorted(set(studies.loc[studies["amr_component"], "study_id"]) - set(amr["study_id"]))
if no_amr_rows:
    report.append("AMR component studies without extracted species-level rows: " + ", ".join(no_amr_rows))
report.append(f"Outcome rows: {len(outcomes)}; driver rows: {len(drivers)}; quality rows: {len(quality)}; discrepancies: {len(discrepancies)}")
report.append(f"Screening records: {len(screening)} (included {int((screening['decision']=='Included').sum())}, excluded {int((screening['decision']=='Excluded').sum())})")
report.append(f"Evidence items resolved verbatim: {len(evidence)}/{len(R.EV_SPECS)}")
report.append("")
report.append(f"Incidence percentage mismatches (> {TOL} points): {len(mismatch_inc)}")
for _, r in mismatch_inc.iterrows():
    report.append(f"  - {r['study_id']} [{r['cohort_label']}]: {r['percent_check']} ({r['n_ssi_events']}/{r['n_total']})")
report.append("")
report.append(f"AMR percentage mismatches (> {TOL} points): {len(mismatch_amr)}")
for _, r in mismatch_amr.iterrows():
    report.append(f"  - {r['study_id']} {r['pathogen']} {r['resistance_outcome']}: {r['percent_check']}")
report.append("")
report.append(f"Derived incidence values: {len(derived_inc)}")
for _, r in derived_inc.iterrows():
    report.append(f"  - {r['study_id']} [{r['cohort_label']}]: {r['notes']}")
report.append("")
report.append(f"Derived AMR values: {len(derived_amr)}")
for _, r in derived_amr.iterrows():
    report.append(f"  - {r['study_id']} {r['pathogen']} {r['resistance_outcome']}: {r['notes']}")
report.append("")
report.append("Priority AMR availability (primary rows):")
priority = amr[amr["include_primary"] == "yes"]
for label, mask in [
    ("MRSA", priority["resistance_outcome"].eq("MRSA")),
    ("3GC/ESBL (E. coli/Klebsiella/Enterobacterales)", priority["resistance_outcome"].isin(["3GC resistance", "3GC non-susceptibility", "ESBL"])),
    ("Carbapenem (A. baumannii/P. aeruginosa/Enterobacterales)", priority["resistance_outcome"].str.startswith("Carbapenem")),
    ("MDR (supplementary, all rows)", amr["resistance_outcome"].eq("MDR")),
]:
    sub = (amr if label.startswith("MDR") else priority)[mask]
    report.append(f"  - {label}: {len(sub)} rows, {sub['study_id'].nunique()} studies")
report.append("")
report.append(f"Structural errors: {len(errors)}")
for e in errors:
    report.append(f"  ! {e}")
report.append("")
report.append("Primary incidence cohorts and denominators:")
for _, r in primary_inc.sort_values(["study_id", "cohort_label"]).iterrows():
    report.append(f"  {r['study_id']:<22} {r['n_ssi_events']:>5}/{r['n_total']:<6} = {r['pct_calculated']}%  {r['percent_check']}")
(OUT / "validation_report.txt").write_text("\n".join(report) + "\n", encoding="utf-8")

# ---------------------------------------------------------------------------
# 8. Excel workbook
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DERIVED_FILL = PatternFill("solid", fgColor="FFF2CC")
MISMATCH_FILL = PatternFill("solid", fgColor="F8CBAD")
SECONDARY_FONT = Font(color="7F7F7F", italic=True)

def xlsx_safe(value):
    """Make corpus text writable by Excel without changing the CSV evidence.

    The PDF corpus contains control characters used as decimal separators
    (\x02), soft hyphens (\xad) and stray control bytes.  CSV keeps the exact
    verbatim slice; the workbook substitutes visible equivalents.
    """
    if not isinstance(value, str):
        return value
    value = value.replace("\x02", ".").replace("\xad", "")
    return "".join(ch if (ch >= " " or ch in "\n\t") else " " for ch in value)


wb = Workbook()
wb.remove(wb.active)

readme = wb.create_sheet("README")
readme_rows = [
    ("ECSA surgical site infection and antimicrobial resistance extraction", ""),
    ("", ""),
    ("Purpose", "Independent, evidence-linked extraction for the systematic review and meta-analysis of SSI burden and AMR patterns in East, Central and Southern Africa."),
    ("Source documents", "56 documents in amr-ssis-ecsa-files; page text extracted with scripts/extract_text.py into corpus/json."),
    ("Evidence model", "Every numeric value links to an evidence_id. The page number and verbatim quote are exported next to the value; quotes are resolved mechanically from the corpus and are never typed into the dataset."),
    ("Primary incidence", "Clinical SSI (CDC definition where reported). Culture-confirmed SSI is recorded separately and never substituted for clinical SSI."),
    ("AMR-only studies", "Studies with an SSI-patient denominator but no operated-patient denominator contribute AMR prevalence only."),
    ("Derived values", "Values calculated from a reported percentage and denominator, or from summed table rows, carry derived_flag='derived' and are explained in notes."),
    ("Arithmetic checks", "percent_check recomputes every percentage from the numerator and denominator. MISMATCH indicates a source inconsistency; the reported counts are retained and the discrepancy is documented."),
    ("Overlapping reports", "Overlap_group marks reports from the same cohort. Duplicates are excluded or merged at study level; clean-cut-ethiopia contains two distinct sequential cohorts."),
    ("Human comparator", "discrepancies.csv documents errors found in the pre-existing human extraction. Human values are not used as a source."),
    ("Reproducibility", "Run: .venv/bin/python scripts/build_dataset.py. Environment: uv, Python 3.12, see pyproject.toml/uv.lock."),
    ("Workbook note", "CSV files contain the exact verbatim quote slices. The XLSX substitutes visible equivalents for PDF control characters (for example the decimal separator control byte becomes a full stop)."),
    ("", ""),
    ("Sheets", "studies, incidence, amr, outcomes, drivers, quality, discrepancies, screening, evidence"),
    ("", ""),
    ("Counts", f"{len(studies)} studies; {len(primary_inc)} primary incidence rows ({primary_inc['study_id'].nunique()} cohorts); {len(amr)} AMR rows; {len(outcomes)} outcome rows; {len(drivers)} driver rows."),
    ("Validation errors", str(len(errors))),
]
for r in readme_rows:
    readme.append(r)
readme.column_dimensions["A"].width = 24
readme.column_dimensions["B"].width = 115
for row in readme.iter_rows(min_row=1, max_row=len(readme_rows)):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
readme["A1"].font = Font(bold=True, size=14, color="1F4E78")

for name, df in tables.items():
    ws = wb.create_sheet(name[:31])
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for _, row in df.iterrows():
        ws.append([None if pd.isna(v) else xlsx_safe(v) for v in row.tolist()])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    # widths and quote wrapping
    for idx, col in enumerate(df.columns, 1):
        letter = get_column_letter(idx)
        if "quote" in col:
            ws.column_dimensions[letter].width = 90
        elif col in {"notes", "reason", "rationale", "issue", "numerator_definition", "denominator_definition",
                     "supports", "title_short", "population", "specialty", "ssi_definition", "denominator_source"}:
            ws.column_dimensions[letter].width = 45
        elif col in {"study_id", "evidence_id", "evidence_id_2", "doc_id", "linked_study_id", "resistance_outcome"}:
            ws.column_dimensions[letter].width = 24
        else:
            ws.column_dimensions[letter].width = 16
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=("quote" in df.columns[cell.column - 1] or
                                                  df.columns[cell.column - 1] in {"notes", "reason", "rationale", "issue"}),
                                       vertical="top")
        # flags
        flags = {df.columns[i - 1]: row[i - 1].value for i in range(1, len(df.columns) + 1)}
        if flags.get("derived_flag") == "derived":
            for cell in row:
                cell.fill = DERIVED_FILL
        if str(flags.get("percent_check", "")).startswith("MISMATCH"):
            for cell in row:
                cell.fill = MISMATCH_FILL
        if flags.get("include_primary") == "no":
            for cell in row:
                cell.font = SECONDARY_FONT

wb.save(OUT / "extraction.xlsx")

# ---------------------------------------------------------------------------
# 9. Console summary
# ---------------------------------------------------------------------------
print("\n".join(report[:14]))
print(f"\nWrote {len(tables)} CSV tables and extraction.xlsx to {OUT}")
if errors:
    print(f"BUILD COMPLETED WITH {len(errors)} STRUCTURAL ERRORS - see validation_report.txt")
    sys.exit(1)
print("Validation: all structural checks passed")
